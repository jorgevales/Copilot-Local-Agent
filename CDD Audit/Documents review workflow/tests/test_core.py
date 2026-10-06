from __future__ import annotations

import csv
import contextlib
import importlib.util
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path


PROJECT_DIR = Path(__file__).resolve().parents[1]
if str(PROJECT_DIR) not in sys.path:
    sys.path.insert(0, str(PROJECT_DIR))

from generate_data import write_analysis_workbooks, write_batch_folders, write_case_csv
from workflow import engine_runner
from workflow.config import WorkflowConfig
from workflow.state import RunState


OPENPYXL_AVAILABLE = importlib.util.find_spec("openpyxl") is not None


def configured(root: Path) -> WorkflowConfig:
    source = root / "source"
    temporary = root / "temporary"
    ready = root / "ready"
    return WorkflowConfig(
        project_root=str(root),
        source_data_root=str(source),
        temporary_batch_dir=str(temporary),
        merged_pdf_dir=str(temporary / "Merged_PDFs"),
        working_csv=str(ready / "working.csv"),
        source_copy_csv=str(ready / "source.csv"),
        completed_csv=str(root / "data" / "completed.csv"),
        sent_log_csv=str(root / "data" / "sent.csv"),
        instructions_md=str(root / "instructions.md"),
        analysis_output_dir=str(root / "analysis"),
        case_size_output_dir=str(root / "sizes"),
        diagnostics_dir=str(root / "diagnostics"),
        edge_profile_dir=str(root / "edge-profile"),
        start_batch=1,
        batch_count=1,
    )


class ConfigAndStateTests(unittest.TestCase):
    def test_config_round_trip_ignores_unknown_fields_and_redacts_profile(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "config.json"
            config = configured(root)
            config.save(path)
            payload = json.loads(path.read_text(encoding="utf-8"))
            payload["unknown_future_setting"] = True
            path.write_text(json.dumps(payload), encoding="utf-8")

            loaded = WorkflowConfig.load(path)
            self.assertEqual(loaded.start_batch, 1)
            self.assertEqual(loaded.path("project_root"), root.resolve())
            self.assertEqual(
                loaded.public_dict()["edge_profile_dir"], "<current-user Edge profile>"
            )

    def test_state_transitions_are_persisted_and_idempotent(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "run" / "state.json"
            state = RunState("synthetic-run")
            state.stage_started("prepare", path)
            state.stage_completed("prepare", path)
            state.stage_completed("prepare", path)
            state.complete(path)

            loaded = RunState.load(path)
            self.assertEqual(loaded.status, "complete")
            self.assertEqual(loaded.completed_stages, ["prepare"])
            self.assertTrue(loaded.finished_at)


class PrepareAdapterTests(unittest.TestCase):
    def test_batch_boundaries(self):
        module = engine_runner.load_file(
            "test_prepare_boundaries",
            engine_runner.REFS
            / "10_prepare_next_ip_batches_long_path_cleanup_sanitized.py",
        )
        self.assertEqual(module.batch_bounds(1), (1, 100))
        self.assertEqual(module.batch_bounds(100), (1, 100))
        self.assertEqual(module.batch_bounds(101), (101, 200))
        plans = module.make_plans(101, 2)
        self.assertEqual([(item.start, item.end) for item in plans], [(101, 200), (201, 300)])

    def test_prepare_copies_100_cases_and_updates_queue_once(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = configured(root)
            write_batch_folders(config.path("source_data_root"), 1)
            write_case_csv(config.path("source_copy_csv"), 1, 100)
            write_case_csv(config.path("working_csv"), 0, 1)

            with contextlib.redirect_stdout(io.StringIO()):
                result = engine_runner.prepare(config, cleanup=False)
            self.assertEqual(result, 0)
            copied = list(config.path("temporary_batch_dir").glob("Change_*"))
            self.assertEqual(len(copied), 100)
            with config.path("working_csv").open(encoding="utf-8", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual(len(rows), 101)
            self.assertEqual(rows[-1]["change_id"], "00100")

            # Recovery/rerun must not duplicate queue rows.
            with contextlib.redirect_stdout(io.StringIO()):
                result = engine_runner.prepare(config, cleanup=False)
            self.assertEqual(result, 0)
            with config.path("working_csv").open(encoding="utf-8", newline="") as handle:
                rerun_rows = list(csv.DictReader(handle))
            self.assertEqual(len(rerun_rows), 101)


@unittest.skipUnless(OPENPYXL_AVAILABLE, "openpyxl is required for workbook tests")
class MasterBatchTests(unittest.TestCase):
    def _master_module(self, folder: Path, name: str):
        module = engine_runner.load_file(
            name,
            engine_runner.REFS
            / "09_append_ip_analysis_to_master_100_case_batches_sanitized.py",
        )
        module.SOURCE_FOLDER = folder
        module.MASTER_PATH = folder / module.MASTER_FILENAME
        module.ERROR_LOG_PATH = folder / module.ERROR_LOG_FILENAME
        return module

    def test_incomplete_batch_is_reported_but_not_ready(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            write_analysis_workbooks(folder, 1, 99)
            module = self._master_module(folder, "test_master_incomplete")
            files = module.select_source_files()
            ready, failed = module.prepare_batches(files)
            self.assertEqual(ready, [])
            self.assertEqual(failed, 0)

    def test_complete_batch_builds_named_master(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config = configured(root)
            write_analysis_workbooks(config.path("analysis_output_dir"), 1, 100)

            self.assertEqual(engine_runner.master(config), 0)
            master = config.path("analysis_output_dir") / (
                "IPs_Completed_Analysis_Master_00001_00100.xlsx"
            )
            self.assertTrue(master.is_file())

            from openpyxl import load_workbook

            workbook = load_workbook(master, read_only=True)
            try:
                self.assertEqual(workbook.active.max_row, 101)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()

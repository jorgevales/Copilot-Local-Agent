from __future__ import annotations

import csv
import os
import re
import sys
import time
import traceback
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ============================================================
# CONFIGURATION
# ============================================================
SOURCE_FOLDER = Path(
    Path.home() / "CopilotCaseAutomation" / "IPs_Documents_Analysis" / "IPs_Completed_Analysis"
)

MASTER_FILENAME_PREFIX = "IPs_Completed_Analysis_Master"
MASTER_FILENAME = f"{MASTER_FILENAME_PREFIX}.xlsx"
ERROR_LOG_FILENAME = "IPs_Completed_Analysis_Errors.csv"

# Fast monitoring and loading settings.
POLL_INTERVAL_SECONDS = 0.50
READ_WORKERS = 16
MAX_SOURCE_OPEN_RETRIES = 20
SOURCE_OPEN_RETRY_SECONDS = 0.25
MAX_MASTER_OPEN_RETRIES = 20
MASTER_OPEN_RETRY_SECONDS = 0.50
CLEANUP_FILE_RETRIES = 20
CLEANUP_FILE_RETRY_SECONDS = 0.25

# Only canonical filenames are appended to the master. Variant copies are cleaned first.
INCOMING_FILE_PATTERN = re.compile(
    r"^Change_[^_]+_IP_[^_]+_documents_analysed\.xlsx$",
    re.IGNORECASE,
)

ANALYSIS_FILE_PATTERN = re.compile(
    r"^(?P<change_label>Change_[^_]+)_IP_(?P<ip_id>[^_]+)_"
    r"documents_analysed(?P<suffix>.*)\.xlsx$",
    re.IGNORECASE,
)

MASTER_PATH = SOURCE_FOLDER / MASTER_FILENAME
ERROR_LOG_PATH = SOURCE_FOLDER / ERROR_LOG_FILENAME


PRIORITY_COLUMNS = [
    "InterestedPartyId",
    "Previous Register Value",
    "Current Register Value",
    "unvalidated_changes",
    "validated_changes",
    "changes_to_documents_relationship",
]

END_COLUMNS = [
    "InterestedPartyCurrentName",
    "ActionDateTime",
    "ActionUserName",
    "ActionUserTeam",
    "ChangedFields",
    "supported_and_unsupported_changes",
    "triggering_documents",
    "documents_quoted_text",
]


# ============================================================
# LOGGING
# ============================================================
def timestamp() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def print_status(message: str) -> None:
    print(f"[{timestamp()}] {message}", flush=True)


def write_error(file_path: Path | None, stage: str, error: BaseException) -> None:
    ERROR_LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
    new_file = not ERROR_LOG_PATH.exists()
    trace = "".join(traceback.format_exception(type(error), error, error.__traceback__))

    with ERROR_LOG_PATH.open("a", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        if new_file:
            writer.writerow([
                "error_datetime", "filename", "full_path", "stage",
                "error_type", "error_message", "traceback",
            ])
        writer.writerow([
            timestamp(),
            file_path.name if file_path else "",
            str(file_path) if file_path else "",
            stage,
            type(error).__name__,
            str(error),
            trace,
        ])
        handle.flush()
        os.fsync(handle.fileno())

    print_status(
        f"ERROR | {file_path.name if file_path else 'General'} | "
        f"{stage} | {type(error).__name__}: {error}"
    )


# ============================================================
# DUPLICATE FILE CLEANUP
# ============================================================
def retry_unlink(path: Path) -> None:
    """Delete a file with retries for short OneDrive or Excel locks."""
    last_error: BaseException | None = None
    for attempt in range(1, CLEANUP_FILE_RETRIES + 1):
        try:
            path.unlink()
            return
        except FileNotFoundError:
            return
        except Exception as error:
            last_error = error
            if attempt < CLEANUP_FILE_RETRIES:
                time.sleep(CLEANUP_FILE_RETRY_SECONDS)
    if last_error is not None:
        raise last_error


def retry_rename(source: Path, target: Path) -> None:
    """Rename a file with retries for short OneDrive or Excel locks."""
    last_error: BaseException | None = None
    for attempt in range(1, CLEANUP_FILE_RETRIES + 1):
        try:
            source.rename(target)
            return
        except Exception as error:
            last_error = error
            if attempt < CLEANUP_FILE_RETRIES:
                time.sleep(CLEANUP_FILE_RETRY_SECONDS)
    if last_error is not None:
        raise last_error


def variant_preference(path: Path) -> tuple[int, int, str]:
    """Rank copies so the highest trailing copy number is kept.

    Examples, from lowest to highest priority:
      - documents_analysed arbitrary.xlsx
      - documents_analysed.xlsx
      - documents_analysed 1.xlsx
      - documents_analysed 2.xlsx
      - documents_analysed 10.xlsx
    """
    match = ANALYSIS_FILE_PATTERN.match(path.name)
    suffix = match.group("suffix").strip() if match else ""

    numeric_match = re.fullmatch(r"(\d+)", suffix)
    if numeric_match:
        return 2, int(numeric_match.group(1)), path.name.lower()

    if suffix == "":
        return 1, 0, path.name.lower()

    return 0, 0, path.name.lower()


def cleanup_duplicate_analysis_files() -> tuple[int, int]:
    """Keep the highest-numbered copy per Change/IP and canonicalise its name."""
    groups: dict[str, dict[str, Any]] = {}
    for path in SOURCE_FOLDER.glob("*.xlsx"):
        if path.name.startswith("~$") or path.name == MASTER_FILENAME:
            continue
        match = ANALYSIS_FILE_PATTERN.match(path.name)
        if not match:
            continue
        canonical_name = (
            f"{match.group('change_label')}_IP_{match.group('ip_id')}_"
            "documents_analysed.xlsx"
        )
        key = canonical_name.lower()
        group = groups.setdefault(key, {
            "canonical_name": canonical_name,
            "change_label": match.group("change_label"),
            "files": [],
        })
        group["files"].append(path)

    deleted_count = 0
    renamed_count = 0
    reports: list[dict[str, Any]] = []

    for key in sorted(groups):
        group = groups[key]
        canonical_name = group["canonical_name"]
        change_label = group["change_label"]
        canonical_path = SOURCE_FOLDER / canonical_name
        files: list[Path] = sorted(
            group["files"],
            key=lambda item: item.name.lower(),
        )

        # max() is intentional: a numbered copy outranks the unnumbered
        # canonical file, and the greatest copy number outranks lower numbers.
        selected_path = max(files, key=variant_preference)
        selected_original_name = selected_path.name
        files_to_delete = [path for path in files if path != selected_path]

        report = {
            "change_label": change_label,
            "kept": canonical_name,
            "selected_original": selected_original_name,
            "renamed_from": None,
            "deleted": [],
        }

        # If a lower-priority canonical file exists, remove it before renaming
        # the selected highest-numbered copy to the canonical filename.
        for duplicate_path in files_to_delete:
            duplicate_name = duplicate_path.name
            retry_unlink(duplicate_path)
            deleted_count += 1
            report["deleted"].append(duplicate_name)

        if selected_path.name.lower() != canonical_name.lower():
            retry_rename(selected_path, canonical_path)
            renamed_count += 1
            report["renamed_from"] = selected_original_name

        if report["renamed_from"] or report["deleted"]:
            reports.append(report)

    if reports:
        print_status("Duplicate cleanup results grouped by Change ID:")
        for report in reports:
            print(f"\n{report['change_label']}", flush=True)
            if report["renamed_from"]:
                print(
                    f"  Selected highest version: {report['renamed_from']}",
                    flush=True,
                )
                print(
                    f"  Renamed: {report['renamed_from']} -> {report['kept']}",
                    flush=True,
                )
            print(f"  Kept: {report['kept']}", flush=True)
            if report["deleted"]:
                print("  Deleted:", flush=True)
                for file_name in report["deleted"]:
                    print(f"    - {file_name}", flush=True)
            else:
                print("  Deleted: None", flush=True)
        print("", flush=True)
    else:
        print_status("Duplicate cleanup: no duplicate files found.")

    return deleted_count, renamed_count

# ============================================================
# FILE DISCOVERY AND READING
# ============================================================
def discover_source_files() -> list[Path]:
    """Return every canonical analysis workbook for this one-shot rebuild."""
    files = [
        path
        for path in SOURCE_FOLDER.glob("*.xlsx")
        if path.name != MASTER_FILENAME
        and not path.name.startswith("~$")
        and INCOMING_FILE_PATTERN.match(path.name)
    ]
    files.sort(key=lambda item: item.name.lower())
    return files


def normalise_header(value: Any) -> str:
    return "" if value is None else str(value).strip()


def cell_value_for_output(cell, header: str) -> Any:
    value = cell.value

    if value is not None and isinstance(value, (int, float)):
        number_format = str(cell.number_format or "")
        if re.fullmatch(r"0+", number_format):
            return f"{int(value):0{len(number_format)}d}"

    if header.lower() == "change_id" and value is not None:
        text = str(value)
        if re.fullmatch(r"\d+(?:\.0+)?", text):
            return str(int(float(text))).zfill(5)
        return text

    return value


def read_source_workbook(file_path: Path) -> tuple[Path, list[str], list[list[Any]], os.stat_result]:
    last_error: BaseException | None = None

    for attempt in range(1, MAX_SOURCE_OPEN_RETRIES + 1):
        workbook = None
        try:
            # A OneDrive placeholder or partial upload can be size 0 temporarily.
            stat_before = file_path.stat()
            if stat_before.st_size <= 0:
                raise OSError("File is empty or still syncing.")

            workbook = load_workbook(
                file_path,
                read_only=True,
                data_only=False,
                keep_links=False,
            )
            worksheet = workbook.active

            header_cells = next(worksheet.iter_rows(min_row=1, max_row=1))
            headers = [normalise_header(cell.value) for cell in header_cells]
            while headers and headers[-1] == "":
                headers.pop()

            if not headers:
                raise ValueError("Row 1 does not contain headers.")
            if any(not header for header in headers):
                raise ValueError("One or more header cells in row 1 are blank.")
            if len(set(headers)) != len(headers):
                raise ValueError("Row 1 contains duplicate headers.")

            rows: list[list[Any]] = []
            for cells in worksheet.iter_rows(min_row=2, max_col=len(headers)):
                values = [
                    cell_value_for_output(cell, headers[index])
                    for index, cell in enumerate(cells)
                ]
                if any(value is not None and str(value).strip() != "" for value in values):
                    rows.append(values)

            workbook.close()
            workbook = None
            return file_path, headers, rows, file_path.stat()

        except Exception as error:
            last_error = error
            if workbook is not None:
                try:
                    workbook.close()
                except Exception:
                    pass
            if attempt < MAX_SOURCE_OPEN_RETRIES:
                time.sleep(SOURCE_OPEN_RETRY_SECONDS)

    if last_error is not None:
        raise last_error
    raise RuntimeError(f"Could not open workbook: {file_path}")


def read_sources_in_parallel(files: list[Path]):
    successful = []
    failed = []
    workers = min(READ_WORKERS, max(len(files), 1))

    with ThreadPoolExecutor(max_workers=workers, thread_name_prefix="excel-reader") as executor:
        futures = {executor.submit(read_source_workbook, path): path for path in files}
        for future in as_completed(futures):
            path = futures[future]
            try:
                successful.append(future.result())
            except Exception as error:
                write_error(path, "read incoming workbook", error)
                failed.append(path)

    successful.sort(
        key=lambda item: (
            int(re.match(r"^Change_(\d+)", item[0].name, re.IGNORECASE).group(1))
            if re.match(r"^Change_(\d+)", item[0].name, re.IGNORECASE)
            else 10**18,
            item[0].name.lower(),
        )
    )
    return successful, failed


# ============================================================
# MASTER WORKBOOK
# ============================================================
def open_master_with_retry():
    last_error: BaseException | None = None
    for attempt in range(1, MAX_MASTER_OPEN_RETRIES + 1):
        try:
            return load_workbook(
                MASTER_PATH,
                read_only=False,
                data_only=False,
                keep_links=False,
            )
        except Exception as error:
            last_error = error
            if attempt < MAX_MASTER_OPEN_RETRIES:
                time.sleep(MASTER_OPEN_RETRY_SECONDS)
    if last_error is not None:
        raise last_error
    raise RuntimeError(f"Could not open master workbook: {MASTER_PATH}")


def create_master_workbook(headers: list[str]):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Combined Analysis"
    worksheet.append(headers)

    header_fill = PatternFill(fill_type="solid", fgColor="000000")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for column_number in range(1, len(headers) + 1):
        cell = worksheet.cell(row=1, column=column_number)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    worksheet.sheet_view.showGridLines = True

    for column_number, header in enumerate(headers, start=1):
        worksheet.column_dimensions[get_column_letter(column_number)].width = min(
            max(len(header) + 2, 12), 35
        )
    return workbook

def save_master_atomically(workbook, master_path: Path) -> None:
    """Save safely; never delete a previous master before the replacement is ready."""
    temporary_path = master_path.with_name(
        f"{master_path.stem}.temporary{master_path.suffix}"
    )
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, master_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()



def update_width_cache(widths: dict[int, int], headers: list[str], rows: list[list[Any]]) -> None:
    for column_number, header in enumerate(headers, start=1):
        longest = widths.get(column_number, len(header))
        for row in rows:
            if column_number > len(row):
                continue
            value = row[column_number - 1]
            if value is None:
                continue
            longest = max(
                longest,
                max((len(line) for line in str(value).splitlines()), default=0),
            )
        widths[column_number] = longest


def preferred_header_order(headers: list[str]) -> list[str]:
    """Place required initial columns first and requested metadata columns last."""
    actual_by_lower = {normalise_header(header).lower(): header for header in headers}
    ordered: list[str] = []
    used: set[str] = set()

    for requested_header in PRIORITY_COLUMNS:
        actual_header = actual_by_lower.get(requested_header.lower())
        if actual_header is not None and actual_header.lower() not in used:
            ordered.append(actual_header)
            used.add(actual_header.lower())

    end_column_names = {header.lower() for header in END_COLUMNS}
    ordered.extend(
        header for header in headers
        if normalise_header(header).lower() not in used
        and normalise_header(header).lower() not in end_column_names
    )
    used.update(normalise_header(header).lower() for header in ordered)

    for requested_header in END_COLUMNS:
        actual_header = actual_by_lower.get(requested_header.lower())
        if actual_header is not None and actual_header.lower() not in used:
            ordered.append(actual_header)
            used.add(actual_header.lower())

    current_name = actual_by_lower.get("interestedpartycurrentname")
    change_id = actual_by_lower.get("change_id")
    if current_name and change_id:
        ordered = [h for h in ordered if normalise_header(h).lower() != "interestedpartycurrentname"]
        pos = next(i for i, h in enumerate(ordered) if normalise_header(h).lower() == "change_id")
        ordered.insert(pos + 1, current_name)
    return ordered


def reorder_rows(
    source_headers: list[str],
    rows: list[list[Any]],
    target_headers: list[str],
) -> list[list[Any]]:
    source_positions = {
        normalise_header(header).lower(): index
        for index, header in enumerate(source_headers)
    }
    reordered_rows = []
    for row in rows:
        reordered_rows.append([
            row[source_positions[target_header.lower()]]
            if source_positions[target_header.lower()] < len(row)
            else None
            for target_header in target_headers
        ])
    return reordered_rows


def format_pipe_separated_values(value: Any) -> Any:
    """Replace one or more pipe separators with one real line break in the same cell."""
    if value is None or not isinstance(value, str) or "|" not in value:
        return value
    return re.sub(r"\s*\|+\s*", "\n", value)


def format_register_value(value: Any) -> Any:
    if value is None or not isinstance(value, str): return value
    text=re.sub(r"(?i)<br\s*/?>","\n",value).replace("\r\n","\n").replace("\r","\n")
    out=[]
    for raw in text.split("\n"):
        line=raw.strip()
        if not line: continue
        addr=bool(re.match(r"(?i)^Address Entry\s+\d+\b",line)); phone=bool(re.match(r"(?i)^Phone Entry\s+\d+\b",line))
        if not (addr or phone): out.append(line); continue
        parts=[x.strip() for x in line.split(';') if x.strip()]; out.append(parts[0]); fields=[]
        for part in parts[1:]:
            key,val=part.split('=',1) if '=' in part else (part,''); fields.append((key.strip(),val.strip(),part))
        if addr:
            vals={}; last=None
            for i,(k,v,_) in enumerate(fields):
                m=re.fullmatch(r"(?i)Address([1-6])",k)
                if m: vals[int(m.group(1))]=v; last=i
            for i,(_,_,part) in enumerate(fields):
                out.append(part)
                if i==last:
                    full=' '.join(vals[n] for n in range(1,7) if vals.get(n))
                    if full: out.append('Full address: '+full)
        else:
            vals={}
            for k,v,part in fields: out.append(part); vals[k.lower()]=v
            req=('internationaldialingcodeid','areacode','phonenumber')
            if all(vals.get(k) for k in req): out.append('Full phone number: '+' '.join(vals[k] for k in req))
    return '\n'.join(out)

UNVALIDATED_BOLD_TEXT=(
 "Unsupported or contradictory changes",
 "SmartSearch conclusion: Possible SmartSearch needed",
 "SmartSearch conclusion: Related SmartSearch found",
)

def smartsearch_conclusions_as_rich_text(value: Any) -> Any:
    if value is None: return value
    text=re.sub(r"(?i)<br\s*/?>","\n",str(value)).replace("\r\n","\n").replace("\r","\n")
    conclusions=[x for x in UNVALIDATED_BOLD_TEXT if x.lower().startswith('smartsearch conclusion:')]
    cp=re.compile('('+'|'.join(re.escape(x) for x in conclusions)+')',re.I)
    text=cp.sub(lambda m:'\n\n'+m.group(1)+'\n\n',text); text=re.sub(r'\n{3,}','\n\n',text).strip()
    bp=re.compile('('+'|'.join(re.escape(x) for x in UNVALIDATED_BOLD_TEXT)+')',re.I)
    rich=CellRichText(); bold=InlineFont(rFont='Calibri',sz=11,b=True); regular=InlineFont(rFont='Calibri',sz=11,b=False); cursor=0; found=False
    for m in bp.finditer(text):
        found=True
        if m.start()>cursor: rich.append(TextBlock(regular,text[cursor:m.start()]))
        rich.append(TextBlock(bold,m.group(1))); cursor=m.end()
    if cursor<len(text): rich.append(TextBlock(regular,text[cursor:]))
    return rich if found else justification_as_rich_text(text)


NUMBERED_CHANGE_PATTERN = re.compile(r"^\s*(\d+)\.\s+(.+?)\s*$")
UNSUPPORTED_CHANGE_PATTERN = re.compile(
    r"^\s*(\d+)\.\s+(.+?)\s+"
    r"\((Partially supported|Not supported|Contradicted)\)\s*$",
    re.IGNORECASE,
)


def normalise_multiline_text(value: Any) -> Any:
    """Normalise cell line endings without changing the visible content."""
    if value is None or not isinstance(value, str):
        return value
    return value.replace("\r\n", "\n").replace("\r", "\n").strip()


def clean_justification_source_text(value: Any) -> str:
    """Convert HTML-style breaks and line endings into clean plain-text lines."""
    if value is None:
        return ""
    text = str(value)
    text = re.sub(r"(?i)<br\s*/?>", "\n", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    return text.strip()


def format_justification_section(title: str, lines: list[str]) -> str:
    """Format one section with its title, then blank lines between numbered items."""
    cleaned_lines = [line.strip() for line in lines if line.strip()]
    if not cleaned_lines:
        return ""

    output = [title]
    for line in cleaned_lines:
        if len(output) > 1 and NUMBERED_CHANGE_PATTERN.match(line):
            output.append("")
        output.append(line)
    return "\n".join(output)


def split_justification(value: Any) -> tuple[str, str]:
    """Split the original justification into unsupported and supported sections."""
    text = clean_justification_source_text(value)
    if not text:
        return "", ""

    sections: dict[str, list[str]] = {"supported": [], "unsupported": []}
    current_section: str | None = None

    for raw_line in text.split("\n"):
        line = raw_line.strip()
        if not line:
            continue

        heading = line.rstrip(":").strip().lower()
        if heading == "supported changes":
            current_section = "supported"
            continue
        if heading in {"unsupported or contradictory changes", "unsupported changes"}:
            current_section = "unsupported"
            continue

        # Preserve unexpected pre-heading text rather than losing it.
        if current_section is None:
            current_section = "supported"
        sections[current_section].append(line)

    unsupported = format_justification_section(
        "Unsupported or contradictory changes",
        sections["unsupported"],
    )
    supported = format_justification_section(
        "Supported changes",
        sections["supported"],
    )
    return unsupported, supported


def split_justification_column(
    headers: list[str],
    rows: list[list[Any]],
) -> tuple[list[str], list[list[Any]]]:
    """Replace the source justification column with the two requested output columns."""
    justification_index = next(
        (
            index
            for index, header in enumerate(headers)
            if normalise_header(header).lower() == "justification"
        ),
        None,
    )
    if justification_index is None:
        return headers, rows

    new_headers = (
        headers[:justification_index]
        + ["unvalidated_changes", "validated_changes"]
        + headers[justification_index + 1:]
    )
    new_rows: list[list[Any]] = []
    for row in rows:
        padded_row = list(row) + [None] * max(0, len(headers) - len(row))
        unsupported, supported = split_justification(padded_row[justification_index])
        new_rows.append(
            padded_row[:justification_index]
            + [unsupported, supported]
            + padded_row[justification_index + 1:]
        )
    return new_headers, new_rows


def justification_as_rich_text(value: Any) -> Any:
    """Make only the section title bold while retaining all requested line breaks."""
    if value is None or not isinstance(value, str) or not value:
        return value

    title, separator, body = value.partition("\n")
    rich_text = CellRichText()
    rich_text.append(TextBlock(InlineFont(rFont="Calibri", sz=11, b=True), title))
    if separator:
        rich_text.append(TextBlock(InlineFont(rFont="Calibri", sz=11, b=False), "\n" + body))
    return rich_text


def format_numbered_relationship(value: Any) -> Any:
    """Keep every numbered relationship row separated by one blank line."""
    if value is None or not isinstance(value, str): return value
    text=re.sub(r"(?i)<br\s*/?>","\n",value).replace("\r\n","\n").replace("\r","\n")
    return "\n\n".join(line.strip() for line in text.split("\n") if line.strip())


def format_supported_and_unsupported_changes(value: Any) -> Any:
    """Normalise the two-section numbered support summary used by the new output."""
    normalised = normalise_multiline_text(value)
    if normalised is None or not isinstance(normalised, str) or not normalised:
        return normalised

    raw_lines = [line.strip() for line in normalised.split("\n") if line.strip()]
    output: list[str] = []
    for line in raw_lines:
        heading = line.rstrip(":").strip().lower()
        if heading == "supported changes":
            if output:
                output.append("")
            output.append("Supported changes")
        elif heading == "unsupported changes":
            if output and output[-1] != "":
                output.append("")
            output.append("Unsupported changes")
        else:
            output.append(line)
    return "\n".join(output)


def confidence_as_excel_number(value: Any) -> float | None:
    """Convert confidence percentages to Excel numeric percentages."""
    if value is None or isinstance(value, bool):
        return None

    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text_value = str(value).strip()
        if not text_value:
            return None
        has_percent_sign = text_value.endswith("%")
        if has_percent_sign:
            text_value = text_value[:-1].strip()
        text_value = text_value.replace(",", ".")
        try:
            number = float(text_value)
        except ValueError:
            return None
        if has_percent_sign:
            number /= 100

    if number > 1:
        number /= 100
    return number


def rebuild_sheet_in_column_order(workbook, worksheet, target_headers: list[str]):
    """Rebuild an existing master sheet when its columns are not in the requested order."""
    current_headers = [
        normalise_header(worksheet.cell(1, column).value)
        for column in range(1, worksheet.max_column + 1)
    ]
    current_headers = [
        "supported_and_unsupported_changes"
        if header.lower() == "supported_changes"
        else header
        for header in current_headers
    ]
    for column_number, header in enumerate(current_headers, start=1):
        worksheet.cell(row=1, column=column_number).value = header
    while current_headers and current_headers[-1] == "":
        current_headers.pop()

    if current_headers == target_headers:
        return worksheet

    if {header.lower() for header in current_headers} != {
        header.lower() for header in target_headers
    }:
        raise ValueError(
            "Master header mismatch.\n"
            f"Required columns: {target_headers}\n"
            f"Master columns: {current_headers}"
        )

    existing_rows = [
        [worksheet.cell(row=row_number, column=column_number).value
         for column_number in range(1, len(current_headers) + 1)]
        for row_number in range(2, worksheet.max_row + 1)
    ]
    reordered_rows = reorder_rows(current_headers, existing_rows, target_headers)

    sheet_index = workbook.worksheets.index(worksheet)
    original_title = worksheet.title
    workbook.remove(worksheet)
    worksheet = workbook.create_sheet(title=original_title, index=sheet_index)
    worksheet.append(target_headers)
    for row in reordered_rows:
        worksheet.append(row)
    worksheet.sheet_view.showGridLines = True
    return worksheet


def apply_master_formatting(worksheet, headers: list[str]) -> None:
    """Apply content conversions and final Excel formatting to the master sheet."""
    header_fill = PatternFill(fill_type="solid", fgColor="000000")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="Calibri", size=11, bold=False, color="000000")
    green_fill = PatternFill(fill_type="solid", fgColor="D8E4BC")
    yellow_fill = PatternFill(fill_type="solid", fgColor="FFFCC9")
    red_fill = PatternFill(fill_type="solid", fgColor="F2DCDB")
    no_fill = PatternFill(fill_type=None)

    header_lookup = {
        normalise_header(header).lower(): index
        for index, header in enumerate(headers, start=1)
    }

    for column_number in range(1, len(headers) + 1):
        cell = worksheet.cell(row=1, column=column_number)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    previous_values_column = header_lookup.get("previous register value")
    new_values_column = header_lookup.get("current register value")
    relationship_column = header_lookup.get("changes_to_documents_relationship")
    support_summary_column = header_lookup.get("supported_and_unsupported_changes")
    unvalidated_changes_column = header_lookup.get("unvalidated_changes")
    validated_changes_column = header_lookup.get("validated_changes")
    confidence_column = header_lookup.get("confidence_level")

    for row_number in range(2, worksheet.max_row + 1):
        if previous_values_column is not None:
            cell = worksheet.cell(row=row_number, column=previous_values_column)
            cell.value = format_pipe_separated_values(cell.value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if new_values_column is not None:
            cell = worksheet.cell(row=row_number, column=new_values_column)
            cell.value = format_register_value(format_pipe_separated_values(cell.value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        if relationship_column is not None:
            cell = worksheet.cell(row=row_number, column=relationship_column)
            cell.value = format_numbered_relationship(cell.value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        if support_summary_column is not None:
            cell = worksheet.cell(row=row_number, column=support_summary_column)
            cell.value = format_supported_and_unsupported_changes(cell.value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        if unvalidated_changes_column is not None:
            cell = worksheet.cell(row=row_number, column=unvalidated_changes_column)
            original_text = str(cell.value) if cell.value is not None else ""
            cell.value = smartsearch_conclusions_as_rich_text(cell.value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            # Apply mutually exclusive background colours to unvalidated_changes.
            normalised_fill_text = re.sub(
                r"(?i)<br\s*/?>",
                "\n",
                original_text,
            )
            normalised_fill_text = normalised_fill_text.replace("\r\n", "\n").replace("\r", "\n")
            normalised_fill_text = "\n".join(
                line.strip()
                for line in normalised_fill_text.split("\n")
                if line.strip()
            )

            possible_smartsearch_text = (
                "SmartSearch conclusion: Possible SmartSearch needed"
            )
            no_unsupported_changes_text = (
                "Unsupported or contradictory changes\nNone"
            )

            if possible_smartsearch_text.lower() in normalised_fill_text.lower():
                cell.fill = PatternFill(fill_type="solid", fgColor="EBF7F9")
            elif normalised_fill_text.lower() == no_unsupported_changes_text.lower():
                cell.fill = PatternFill(fill_type=None)
            else:
                cell.fill = PatternFill(fill_type="solid", fgColor="F3F2E9")
        if validated_changes_column is not None:
            cell = worksheet.cell(row=row_number, column=validated_changes_column)
            cell.value = justification_as_rich_text(cell.value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        if confidence_column is not None:
            cell = worksheet.cell(row=row_number, column=confidence_column)
            numeric_value = confidence_as_excel_number(cell.value)
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)

            if numeric_value is None:
                cell.fill = no_fill
            else:
                cell.value = numeric_value
                cell.number_format = "0%"
                if 0.70 <= numeric_value <= 1.00:
                    cell.fill = green_fill
                elif 0.50 <= numeric_value < 0.70:
                    cell.fill = yellow_fill
                elif numeric_value < 0.50:
                    cell.fill = red_fill
                else:
                    cell.fill = no_fill


def build_master_workbook(successful_files, master_path: Path) -> int:
    """Build one complete batch using the original layout and formatting."""
    if not successful_files:
        return 0

    transformed_files = []
    for file_path, headers, rows, stat in successful_files:
        split_headers, split_rows = split_justification_column(headers, rows)
        renamed_headers = ["Previous Register Value" if normalise_header(h).lower()=="previousvalues" else "Current Register Value" if normalise_header(h).lower()=="newvalues" else h for h in split_headers]
        transformed_files.append((file_path, renamed_headers, split_rows, stat))

    source_headers = transformed_files[0][1]
    expected_headers = preferred_header_order(source_headers)
    valid_files = []

    for file_path, headers, rows, stat in transformed_files:
        if headers != source_headers:
            error = ValueError(
                f"Header mismatch in {file_path.name}. "
                f"Expected {source_headers}; found {headers}"
            )
            write_error(file_path, "validate headers", error)
        else:
            reordered_rows = reorder_rows(headers, rows, expected_headers)
            valid_files.append((file_path, expected_headers, reordered_rows, stat))

    if not valid_files:
        return 0

    workbook = create_master_workbook(expected_headers)

    try:
        worksheet = workbook.active
        worksheet = rebuild_sheet_in_column_order(workbook, worksheet, expected_headers)
        workbook.active = workbook.worksheets.index(worksheet)

        widths = {
            column_number: int(
                worksheet.column_dimensions[get_column_letter(column_number)].width or 10
            )
            for column_number in range(1, len(expected_headers) + 1)
        }
        change_id_column = next(
            (
                index
                for index, header in enumerate(expected_headers, start=1)
                if header.lower() == "change_id"
            ),
            None,
        )

        total_rows_added = 0
        completion_data = []
        first_new_row = worksheet.max_row + 1

        for file_path, headers, rows, stat in valid_files:
            for row in rows:
                worksheet.append(row)
            update_width_cache(widths, headers, rows)
            total_rows_added += len(rows)
            completion_data.append((file_path, len(rows), stat))

        last_new_row = worksheet.max_row
        if last_new_row >= first_new_row:
            for row in worksheet.iter_rows(
                min_row=first_new_row,
                max_row=last_new_row,
                max_col=len(expected_headers),
            ):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                if change_id_column is not None:
                    row[change_id_column - 1].number_format = "@"

        apply_master_formatting(worksheet, expected_headers)

        long_text_headers = {
            "previous register value", "current register value", "triggering_documents",
            "documents_quoted_text", "changes_to_documents_relationship",
            "supported_and_unsupported_changes",
            "unvalidated_changes", "validated_changes",
            "potential_error",
        }
        for column_number, header in enumerate(expected_headers, start=1):
            maximum = 60 if header.lower() in long_text_headers else 35
            worksheet.column_dimensions[get_column_letter(column_number)].width = min(
                max(widths.get(column_number, len(header)) + 2, 10), maximum
            )

        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.freeze_panes = "A2"
        for column_number in range(9, 24):
            worksheet.column_dimensions[get_column_letter(column_number)].hidden = True

        save_master_atomically(workbook, master_path)

        master_rows = max(worksheet.max_row - 1, 0)
        for file_path, rows_added, stat in completion_data:
            print_status(
                f"Completed: {file_path.name} | Rows added: {rows_added}"
            )

        print_status(
            f"Batch saved | {master_path.name} | Files: {len(completion_data)} | "
            f"Rows appended: {total_rows_added} | "
            f"Master rows excluding header: {master_rows}"
        )
        return len(completion_data)

    finally:
        workbook.close()


# ============================================================
# CHANGE ID SEQUENCE REPORTING
# ============================================================
CHANGE_ID_PATTERN = re.compile(r"^Change_(?P<change_id>\d+)$", re.IGNORECASE)

def configure_master_path_for_files(files: list[Path]) -> Path:
    """Name the master from the lowest and highest Change IDs in this run."""
    global MASTER_FILENAME, MASTER_PATH

    numeric_ids: list[int] = []
    source_widths: list[int] = []
    for file_path in files:
        analysis_match = ANALYSIS_FILE_PATTERN.match(file_path.name)
        if analysis_match is None:
            continue
        change_match = CHANGE_ID_PATTERN.fullmatch(analysis_match.group("change_label"))
        if change_match is None:
            continue
        raw_change_id = change_match.group("change_id")
        numeric_ids.append(int(raw_change_id))
        source_widths.append(len(raw_change_id))

    if not numeric_ids:
        raise ValueError(
            "Cannot name the master workbook because no numeric Change IDs were found."
        )

    display_width = max(5, max(source_widths, default=0))
    first_change_id = min(numeric_ids)
    last_change_id = max(numeric_ids)
    MASTER_FILENAME = (
        f"{MASTER_FILENAME_PREFIX}_{first_change_id:0{display_width}d}_"
        f"{last_change_id:0{display_width}d}.xlsx"
    )
    MASTER_PATH = SOURCE_FOLDER / MASTER_FILENAME
    return MASTER_PATH


def collapse_consecutive_ids(change_ids: list[int]) -> list[tuple[int, int]]:
    """Collapse sorted unique numeric Change IDs into inclusive ranges."""
    if not change_ids:
        return []

    ranges: list[tuple[int, int]] = []
    start = previous = change_ids[0]

    for change_id in change_ids[1:]:
        if change_id == previous + 1:
            previous = change_id
            continue

        ranges.append((start, previous))
        start = previous = change_id

    ranges.append((start, previous))
    return ranges


def format_change_id_range(start: int, end: int, width: int) -> str:
    """Format a single Change ID or an inclusive Change ID range."""
    if start == end:
        return f"Change ID {start:0{width}d}"
    return f"Change ID {start:0{width}d} - Change ID {end:0{width}d}"


def report_change_id_sequences(files: list[Path]) -> None:
    """Print present consecutive Change ID ranges and gaps between min/max IDs."""
    numeric_ids: set[int] = set()
    source_widths: list[int] = []
    unrecognised_files: list[str] = []

    for file_path in files:
        analysis_match = ANALYSIS_FILE_PATTERN.match(file_path.name)
        if analysis_match is None:
            unrecognised_files.append(file_path.name)
            continue

        change_label = analysis_match.group("change_label")
        change_match = CHANGE_ID_PATTERN.fullmatch(change_label)
        if change_match is None:
            unrecognised_files.append(file_path.name)
            continue

        raw_change_id = change_match.group("change_id")
        numeric_ids.add(int(raw_change_id))
        source_widths.append(len(raw_change_id))

    print()
    print("Consecutive IDs:")

    if not numeric_ids:
        print("None found")
        print()
        print("Missing Change IDs:")
        print("Cannot determine missing IDs because no numeric Change IDs were found.")
    else:
        ordered_ids = sorted(numeric_ids)
        display_width = max(5, max(source_widths, default=0))

        for start, end in collapse_consecutive_ids(ordered_ids):
            print(format_change_id_range(start, end, display_width))

        present_ids = set(ordered_ids)
        missing_ids = [
            change_id
            for change_id in range(ordered_ids[0], ordered_ids[-1] + 1)
            if change_id not in present_ids
        ]

        print()
        print("Missing Change IDs:")
        if missing_ids:
            for start, end in collapse_consecutive_ids(missing_ids):
                print(format_change_id_range(start, end, display_width))
        else:
            print("None")

    if unrecognised_files:
        print()
        print("Files excluded from the Change ID sequence report:")
        for file_name in sorted(unrecognised_files, key=str.lower):
            print(file_name)

    print()


# ============================================================
# ONE-SHOT EXECUTION
# ============================================================
# ============================================================
# COMPLETE, FIXED 100-CHANGE BATCHES
# ============================================================
BATCH_SIZE = 100


def batch_start(change_id: int) -> int:
    """IDs 1-100, 101-200, 5001-5100, etc."""
    if change_id < 1:
        raise ValueError("Change IDs must be positive integers.")
    return ((change_id - 1) // BATCH_SIZE) * BATCH_SIZE + 1


def source_change_id(path: Path) -> int | None:
    match = ANALYSIS_FILE_PATTERN.fullmatch(path.name)
    if not match:
        return None
    label = CHANGE_ID_PATTERN.fullmatch(match.group("change_label"))
    if not label:
        return None
    value = int(label.group("change_id"))
    return value if value > 0 else None


def master_path_for_batch(start: int, width: int) -> Path:
    end = start + BATCH_SIZE - 1
    return SOURCE_FOLDER / (
        f"{MASTER_FILENAME_PREFIX}_{start:0{width}d}_{end:0{width}d}.xlsx"
    )


def select_source_files() -> list[Path]:
    """Choose the same highest-numbered copy as the original cleanup, without deleting files."""
    chosen: dict[str, Path] = {}
    for path in SOURCE_FOLDER.glob("*.xlsx"):
        if path.name.startswith("~$") or source_change_id(path) is None:
            continue
        match = ANALYSIS_FILE_PATTERN.fullmatch(path.name)
        key = (f"{match.group('change_label')}_IP_{match.group('ip_id')}_"
               "documents_analysed.xlsx").lower()
        if key not in chosen or variant_preference(path) > variant_preference(chosen[key]):
            chosen[key] = path
    return sorted(chosen.values(), key=lambda path: (source_change_id(path), path.name.lower()))


def prepare_batches(files: list[Path]):
    """Read sources and qualify each batch independently before offering creation."""
    report_change_id_sequences(files)
    width = max(5, *(len(CHANGE_ID_PATTERN.fullmatch(
        ANALYSIS_FILE_PATTERN.fullmatch(p.name).group("change_label")
    ).group("change_id")) for p in files))
    print_status(f"Found {len(files)} selected source workbook(s). Checking readability...")
    successful, failed = read_sources_in_parallel(files)
    by_id: dict[int, list] = {}
    for item in successful:
        by_id.setdefault(source_change_id(item[0]), []).append(item)
    all_ids = {source_change_id(p) for p in files}
    starts = range(batch_start(min(all_ids)), batch_start(max(all_ids)) + 1, BATCH_SIZE)
    ready = []
    print("\nBatch assessment:")
    for start in starts:
        end = start + BATCH_SIZE - 1
        path = master_path_for_batch(start, width)
        missing = [i for i in range(start, end + 1) if i not in all_ids]
        unreadable = [i for i in range(start, end + 1) if i in all_ids and i not in by_id]
        items = [item for i in range(start, end + 1) for item in by_id.get(i, [])]
        problems = []
        if missing:
            problems.append(f"missing IDs: {', '.join(f'{i:0{width}d}' for i in missing)}")
        if unreadable:
            problems.append(f"unreadable IDs: {', '.join(f'{i:0{width}d}' for i in unreadable)}")
        if items:
            baseline = None
            for file_path, headers, rows, stat in items:
                transformed, _ = split_justification_column(headers, rows)
                transformed = ["Previous Register Value" if normalise_header(h).lower() == "previousvalues"
                               else "Current Register Value" if normalise_header(h).lower() == "newvalues"
                               else h for h in transformed]
                if baseline is None:
                    baseline = transformed
                if transformed != baseline:
                    problems.append(f"header mismatch: {file_path.name}")
                if not rows:
                    problems.append(f"no data rows: {file_path.name}")
                id_column = next((j for j, h in enumerate(headers)
                                  if h.lower() == "change_id"), None)
                if id_column is not None:
                    expected = source_change_id(file_path)
                    if any(str(row[id_column]).strip().lstrip('0') != str(expected)
                           for row in rows if row[id_column] is not None
                           and str(row[id_column]).strip()):
                        problems.append(f"Change_ID does not match filename: {file_path.name}")
        if problems:
            print(f"  {start:0{width}d}-{end:0{width}d}: NOT CREATED | " + "; ".join(problems))
        else:
            print(f"  {start:0{width}d}-{end:0{width}d}: READY | {path.name} | {len(items)} source file(s)" +
                  (" | replaces existing master" if path.exists() else ""))
            ready.append((path, items))
    print(f"\nMaster Excel files available to create: {len(ready)}")
    for path, _ in ready:
        print(f"  {path.name}")
    print()
    return ready, len(failed)


def main() -> None:
    print_status("Starting complete 100-change IP analysis master-file build.")
    print_status(f"Source folder: {SOURCE_FOLDER}")
    if not SOURCE_FOLDER.is_dir():
        raise FileNotFoundError(f"Source folder does not exist: {SOURCE_FOLDER}")
    files = select_source_files()
    if not files:
        print_status("No numeric Change ID source workbooks found. No masters created.")
        return
    ready, failed = prepare_batches(files)
    if not ready:
        print_status("No complete, readable 100-change batches. No masters created.")
        return
    try:
        confirmation = input("Press Enter to create the listed master files, or type anything to cancel: ")
    except EOFError:
        print_status("No interactive confirmation received. No masters created.")
        return
    if confirmation != "":
        print_status("Cancelled. No masters created.")
        return
    created = 0
    for path, items in ready:
        try:
            # A source can change while confirmation is pending. Recheck this batch.
            fresh, errors = read_sources_in_parallel([item[0] for item in items])
            failed += len(errors)
            if len(fresh) != len(items) or len({source_change_id(item[0]) for item in fresh}) != BATCH_SIZE:
                print_status(f"Skipped {path.name}: source files changed or could not be read.")
                continue
            baseline = None
            valid = True
            for file_path, headers, rows, stat in fresh:
                transformed, _ = split_justification_column(headers, rows)
                transformed = ["Previous Register Value" if h.lower() == "previousvalues"
                               else "Current Register Value" if h.lower() == "newvalues"
                               else h for h in transformed]
                if baseline is None:
                    baseline = transformed
                if transformed != baseline or not rows:
                    valid = False
                id_column = next((j for j, h in enumerate(headers) if h.lower() == "change_id"), None)
                if id_column is not None and any(
                    str(row[id_column]).strip().lstrip('0') != str(source_change_id(file_path))
                    for row in rows if row[id_column] is not None and str(row[id_column]).strip()
                ):
                    valid = False
            if not valid:
                print_status(f"Skipped {path.name}: source content changed since preview.")
                continue
            if build_master_workbook(fresh, path) == len(fresh):
                created += 1
        except Exception as error:
            write_error(path, "build and save batch master", error)
            failed += 1
    print_status(f"Finished | Master files created: {created}/{len(ready)} | Read/build failures: {failed}")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print_status("Stopped by user.")
    except Exception as fatal_error:
        try:
            write_error(None, "fatal startup error", fatal_error)
        except Exception:
            print(f"Fatal error: {fatal_error}", file=sys.stderr, flush=True)
        sys.exit(1)
